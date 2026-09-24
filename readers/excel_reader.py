
from pathlib import Path
from openpyxl import load_workbook


def read_excel_requests(file_path):

    file_path = Path(file_path)

    if not file_path.is_file():
        raise FileNotFoundError(
            f"Excel template not found: {file_path}"
        )

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )

    try:
        sheet = workbook["Test_Cases"]

        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)

        test_requests = []

        for row in rows:

            if not any(value is not None for value in row):
                continue

            test = dict(zip(headers, row))

            if str(test["Enabled"]).strip().upper() != "YES":
                continue

            request = {
                "test_case_id": test["Test_Case_ID"],
                "environment": test["Environment"],
                "load_plan": test["Load_Plan"],
                "source_system": test["Source_System"],
                "target_system": test["Target_System"],
                "load_type": test["Load_Type"],
            }

            test_requests.append(request)

        return test_requests

    finally:
        workbook.close()
